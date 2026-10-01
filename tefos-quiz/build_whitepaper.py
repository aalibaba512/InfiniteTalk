#!/usr/bin/env python3
"""
Converts WHITEPAPER.md into a formatted .docx, so the paper can be circulated
the way a school committee actually circulates something.
"""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = "/home/user/InfiniteTalk/tefos-quiz"
SRC = f"{HERE}/WHITEPAPER.md"
OUT = f"{HERE}/TEFOS-Quiz-Whitepaper.docx"

TEAL  = RGBColor(0x0B, 0x6B, 0x7D)
INK   = RGBColor(0x16, 0x23, 0x2B)
GREY  = RGBColor(0x5C, 0x6B, 0x75)

doc = Document()

# page + base style
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(0.9)
base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10.5)
base.font.color.rgb = INK
base.paragraph_format.space_after = Pt(7)
base.paragraph_format.line_spacing = 1.18

def add_runs(par, text):
    """Inline **bold**, *italic*, `code`."""
    for tok in re.split(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            par.add_run(tok)
    return par

def style_heading(par, level):
    sizes = {1: 18, 2: 14, 3: 11.5, 4: 10.5}
    par.runs[0].font.size = Pt(sizes.get(level, 10.5))
    par.runs[0].font.color.rgb = TEAL if level <= 2 else INK
    par.runs[0].bold = True
    par.paragraph_format.space_before = Pt(14 if level == 1 else 11)
    par.paragraph_format.space_after = Pt(5)
    return par

lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
first_h1 = True
while i < len(lines):
    ln = lines[i]
    stripped = ln.strip()

    if not stripped:
        i += 1; continue

    # table: a row starting with | followed by a |---| separator
    if stripped.startswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i].strip()); i += 1
        rows = []
        for r in block:
            cells = [c.strip() for c in r.strip("|").split("|")]
            if all(set(c) <= set("-: ") and c for c in cells):
                continue                      # separator row
            rows.append(cells)
        if not rows:
            continue
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncol)
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, r in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncol):
                val = r[ci] if ci < len(r) else ""
                p = cells[ci].paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, val)
                for run in p.runs:
                    run.font.size = Pt(9)
                    if ri == 0:
                        run.bold = True
        doc.add_paragraph()
        continue

    # horizontal rule
    if stripped in ("---", "***", "___"):
        i += 1; continue

    # headings
    m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if m:
        level = len(m.group(1))
        txt = m.group(2)
        if level == 1 and first_h1:
            p = doc.add_paragraph()
            r = p.add_run(txt)
            r.font.size = Pt(24); r.bold = True; r.font.color.rgb = TEAL
            p.paragraph_format.space_after = Pt(4)
            first_h1 = False
        else:
            p = doc.add_paragraph()
            p.add_run(txt)
            style_heading(p, level)
        i += 1; continue

    # blockquote
    if stripped.startswith("> "):
        p = doc.add_paragraph()
        add_runs(p, stripped[2:])
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.space_before = Pt(6)
        for run in p.runs:
            run.italic = True
            run.font.color.rgb = TEAL
        i += 1; continue

    # bullets
    m = re.match(r"^(\s*)[-*]\s+(.*)$", ln)
    if m:
        indent = len(m.group(1)) // 2
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28 + 0.25 * indent)
        p.paragraph_format.space_after = Pt(3)
        add_runs(p, m.group(2))
        i += 1; continue

    # ordered list
    m = re.match(r"^(\s*)\d+\.\s+(.*)$", ln)
    if m:
        indent = len(m.group(1)) // 3
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.28 + 0.25 * indent)
        p.paragraph_format.space_after = Pt(3)
        add_runs(p, m.group(2))
        i += 1; continue

    # paragraph (join soft-wrapped lines)
    para = [stripped]
    i += 1
    while i < len(lines):
        nxt = lines[i]
        s2 = nxt.strip()
        if (not s2 or s2.startswith(("#", "|", ">", "- ", "---"))
                or re.match(r"^\s*\d+\.\s", nxt)):
            break
        para.append(s2)
        i += 1
    p = doc.add_paragraph()
    add_runs(p, " ".join(para))

doc.save(OUT)
print(f"Saved {OUT}")
