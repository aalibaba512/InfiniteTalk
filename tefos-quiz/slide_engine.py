#!/usr/bin/env python3
"""
Slide design & rendering engine for the Telecom Foundation TEFOS Quiz Proposal.
Renders 22 widescreen slides (1920x1080 / 16:9) into:
  1. High-resolution vector/raster PDF (TEFOS-Quiz-Proposal.pdf)
  2. Editable PowerPoint presentation (TEFOS-Quiz-Proposal.pptx)
  3. Individual HD PNG slides (tefos-quiz/slides/slide_01.png .. slide_22.png)
"""
import os
import re
import io
import base64
import html
import pymupdf
import resvg_py
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FONT_DIR = os.path.join(ROOT, "maps", "pakistan", "design", "fonts")
DEJAVU_DIR = "/usr/share/fonts/truetype/dejavu"
ASSETS_DIR = os.path.join(HERE, "assets")
TWEMOJI_DIR = os.path.join(ASSETS_DIR, "twemoji")
LUCIDE_DIR = os.path.join(ASSETS_DIR, "lucide")
PROC_DIR = os.path.join(ASSETS_DIR, "processed")
SLIDES_OUT_DIR = os.path.join(HERE, "slides")

os.makedirs(SLIDES_OUT_DIR, exist_ok=True)

# ---------------------------------------------------------------------------
# Telecom Foundation Brand Palette & Harmonious Extended System
# ---------------------------------------------------------------------------
PAL = {
    # Primary Brand Greens (from official Telecom Foundation logo)
    "TF_DEEP":      "#062B1C",   # Deep Forest Emerald (dark slide canvas)
    "TF_DARK":      "#0A3B26",   # Elevated Dark Emerald Card
    "TF_DARK_CARD": "#0E472F",   # Glass card on dark background
    "TF_DARK_LINE": "#1D6344",   # Border on dark background
    "TF_EMERALD":   "#008450",   # Official 'TELECOM FOUNDATION' logotype green
    "TF_EMERALD_DK":"#00683E",   # Darker emerald for hover/accents
    "TF_LIME":      "#7CC040",   # Official emblem & 'Transforming Communities' lime
    "TF_LIME_BRT":  "#8CE04A",   # Luminous lime for dark backgrounds
    "TF_LIME_PALE": "#E6F5D8",   # Soft lime pill background
    "TF_MINT":      "#EBF6EE",   # Soft emerald-mint container tint
    "TF_MINT_BDR":  "#BFE0C9",   # Mint container border

    # Neutrals & Warm Canvas (inspired by Quetta before & after.pdf)
    "CANVAS":       "#F5F7F2",   # Warm ivory-sage canvas
    "CANVAS_ALT":   "#EEF3EE",   # Subtle secondary row fill
    "WHITE":        "#FFFFFF",   # Pure white card surface
    "BORDER":       "#DCE6DF",   # Crisp sage-tinted card border
    "INK":          "#0F241A",   # Deep forest charcoal for primary text
    "SLATE":        "#486054",   # Muted sage-slate for secondary body copy
    "MUTED":        "#7A9285",   # Footers & captions on light slides
    "DARK_TEXT":    "#E8F5ED",   # Secondary text on dark slides
    "DARK_MUTED":   "#99BDA9",   # Muted text on dark slides

    # Semantic Harmonised Accents
    "GOLD":         "#D99621",   # Rich championship gold (for light backgrounds)
    "GOLD_BRT":     "#F6C453",   # Luminous gold (for dark backgrounds)
    "GOLD_BG":      "#FFF8E6",   # Soft warm amber/gold tint
    "GOLD_BDR":     "#F0D699",   # Gold border
    "CORAL":        "#C93838",   # Alert / Anti-pattern red
    "CORAL_BRT":    "#FF7A7A",   # Coral on dark backgrounds
    "CORAL_BG":     "#FDF0F0",   # Soft coral tint
    "CORAL_BDR":    "#F2C2C2",   # Coral border
}

# ---------------------------------------------------------------------------
# Font Loading & Exact Pixel Measurement via PIL
# ---------------------------------------------------------------------------
_FONT_CACHE = {}

def get_pil_font(family="Poppins", weight="regular", size_px=24):
    key = (family, weight, int(round(size_px)))
    if key not in _FONT_CACHE:
        if family == "Bebas Neue":
            path = os.path.join(FONT_DIR, "BebasNeue-Regular.ttf")
        elif family == "DejaVu Sans Mono":
            fn = "DejaVuSansMono-Bold.ttf" if weight in ("bold", "extrabold", "700", "800") else "DejaVuSansMono.ttf"
            path = os.path.join(DEJAVU_DIR, fn)
        else:
            wmap = {
                "regular": "Poppins-Regular.ttf",
                "400": "Poppins-Regular.ttf",
                "medium": "Poppins-Medium.ttf",
                "500": "Poppins-Medium.ttf",
                "semibold": "Poppins-Bold.ttf",
                "600": "Poppins-Bold.ttf",
                "bold": "Poppins-Bold.ttf",
                "700": "Poppins-Bold.ttf",
                "extrabold": "Poppins-ExtraBold.ttf",
                "800": "Poppins-ExtraBold.ttf",
            }
            path = os.path.join(FONT_DIR, wmap.get(str(weight).lower(), "Poppins-Regular.ttf"))
        _FONT_CACHE[key] = ImageFont.truetype(path, int(round(size_px)))
    return _FONT_CACHE[key]

def measure_text_px(text_str, family="Poppins", weight="regular", size_px=24, letter_spacing=0.0):
    if not text_str:
        return 0.0
    f = get_pil_font(family, weight, size_px)
    w = f.getlength(text_str)
    if letter_spacing and len(text_str) > 1:
        w += letter_spacing * (len(text_str) - 1)
    return w

# ---------------------------------------------------------------------------
# Asset & Icon Cache
# ---------------------------------------------------------------------------
_B64_CACHE = {}
_SVG_CACHE = {}

def image_b64(rel_path):
    full = rel_path if os.path.isabs(rel_path) else os.path.join(ASSETS_DIR, rel_path)
    if full not in _B64_CACHE:
        with open(full, "rb") as f:
            raw = f.read()
        mime = "image/png" if full.lower().endswith(".png") else "image/jpeg"
        _B64_CACHE[full] = f"data:{mime};base64," + base64.b64encode(raw).decode("ascii")
    return _B64_CACHE[full]

def _svg_inner(s):
    s = re.sub(r'<!--.*?-->', '', s, flags=re.S)
    s = re.sub(r'<\?xml.*?\?>', '', s, flags=re.S)
    s = re.sub(r'<svg\b[^>]*>', '', s, count=1, flags=re.S)
    s = re.sub(r'</svg>\s*$', '', s.strip(), flags=re.S)
    return s.strip()

def load_twemoji_svg(name):
    if name not in _SVG_CACHE:
        path = os.path.join(TWEMOJI_DIR, f"{name}.svg")
        with open(path, "r", encoding="utf-8") as f:
            s = f.read()
        m = re.search(r'viewBox="([^"]+)"', s)
        vb = m.group(1) if m else "0 0 36 36"
        inner = _svg_inner(s)
        _SVG_CACHE[name] = (vb, inner)
    return _SVG_CACHE[name]

def load_lucide_svg(name, stroke_color="#008450", stroke_width=2.2):
    key = ("lucide", name, stroke_color, stroke_width)
    if key not in _SVG_CACHE:
        path = os.path.join(LUCIDE_DIR, f"{name}.svg")
        with open(path, "r", encoding="utf-8") as f:
            s = f.read()
        inner = _svg_inner(s)
        grouped = (
            f'<g fill="none" stroke="{stroke_color}" stroke-width="{stroke_width}" '
            f'stroke-linecap="round" stroke-linejoin="round">{inner}</g>'
        )
        _SVG_CACHE[key] = ("0 0 24 24", grouped)
    return _SVG_CACHE[key]

def hex_to_rgb_tuple(hex_str):
    h = hex_str.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

# ---------------------------------------------------------------------------
# SlideCanvas Class (1920 x 1080 coordinate space; 144 px = 1 inch)
# ---------------------------------------------------------------------------
class SlideCanvas:
    W = 1920
    H = 1080
    M = 96  # Standard horizontal margin (96px = 0.667 in)

    def __init__(self, slide_num, dark=False, count=True, bg_image=None, section_tag=None):
        self.slide_num = slide_num
        self.dark = dark
        self.count = count
        self.bg_image = bg_image
        self.section_tag = section_tag
        self.defs = []
        self.bg_els = []
        self.text_els = []
        self.pptx_texts = []
        self._clip_id = 0

        self._init_defs()
        self._init_background()

    def _next_id(self, prefix="clip"):
        self._clip_id += 1
        return f"s{self.slide_num}_{prefix}_{self._clip_id}"

    def _init_defs(self):
        # Soft elevation shadows & brand gradients
        self.defs.append('''
        <filter id="card_shadow" x="-8%" y="-6%" width="116%" height="118%">
          <feDropShadow dx="0" dy="6" stdDeviation="12" flood-color="#072618" flood-opacity="0.06"/>
        </filter>
        <filter id="card_shadow_lg" x="-10%" y="-8%" width="120%" height="124%">
          <feDropShadow dx="0" dy="10" stdDeviation="20" flood-color="#051E13" flood-opacity="0.14"/>
        </filter>
        <linearGradient id="tf_brand_grad" x1="0%" y1="0%" x2="100%" y2="0%">
          <stop offset="0%" stop-color="#008450"/>
          <stop offset="65%" stop-color="#38A348"/>
          <stop offset="100%" stop-color="#7CC040"/>
        </linearGradient>
        <linearGradient id="tf_dark_grad" x1="0%" y1="0%" x2="100%" y2="100%">
          <stop offset="0%" stop-color="#0B422B"/>
          <stop offset="100%" stop-color="#052316"/>
        </linearGradient>
        <linearGradient id="tf_gold_grad" x1="0%" y1="0%" x2="100%" y2="0%">
          <stop offset="0%" stop-color="#D99621"/>
          <stop offset="100%" stop-color="#F6C453"/>
        </linearGradient>
        <radialGradient id="ambient_lime" cx="88%" cy="12%" r="48%">
          <stop offset="0%" stop-color="#7CC040" stop-opacity="0.11"/>
          <stop offset="100%" stop-color="#7CC040" stop-opacity="0.0"/>
        </radialGradient>
        <radialGradient id="ambient_emerald" cx="10%" cy="92%" r="45%">
          <stop offset="0%" stop-color="#008450" stop-opacity="0.07"/>
          <stop offset="100%" stop-color="#008450" stop-opacity="0.0"/>
        </radialGradient>
        ''')

    def _init_background(self):
        base_col = PAL["TF_DEEP"] if self.dark else PAL["CANVAS"]
        self.bg_els.append(f'<rect x="0" y="0" width="{self.W}" height="{self.H}" fill="{base_col}"/>')
        if self.bg_image:
            b64 = image_b64(self.bg_image)
            self.bg_els.append(
                f'<image x="0" y="0" width="{self.W}" height="{self.H}" '
                f'preserveAspectRatio="none" href="{b64}"/>'
            )
        elif not self.dark:
            # Subtle ambient brand glow & geometric wave lines in top-right corner
            self.bg_els.append(f'<rect x="0" y="0" width="{self.W}" height="{self.H}" fill="url(#ambient_lime)"/>')
            self.bg_els.append(f'<rect x="0" y="0" width="{self.W}" height="{self.H}" fill="url(#ambient_emerald)"/>')
            # Delicate abstract geometric rings in top right
            self.bg_els.append(
                '<g opacity="0.22" fill="none" stroke="#7CC040" stroke-width="1.5">'
                '<circle cx="1820" cy="110" r="180"/>'
                '<circle cx="1820" cy="110" r="240" stroke="#008450" stroke-dasharray="6 8"/>'
                '<circle cx="1820" cy="110" r="300"/>'
                '</g>'
            )
        else:
            # Subtle abstract geometry on dark slides without full-bleed photo
            self.bg_els.append(
                '<g opacity="0.16" fill="none" stroke="#8CE04A" stroke-width="1.5">'
                '<circle cx="1780" cy="160" r="220"/>'
                '<circle cx="1780" cy="160" r="310" stroke-dasharray="8 10"/>'
                '</g>'
            )

    def add_header(self, section_label="TEFOS QUIZ PROPOSAL"):
        """Adds the top Telecom Foundation brand bar, logo, and slide counter."""
        # Top gradient accent strip (6px tall)
        self.bg_els.append(f'<rect x="0" y="0" width="{self.W}" height="6" fill="url(#tf_brand_grad)"/>')

        # Official Telecom Foundation Logo (top-left)
        logo_file = "tf_logo_dark.png" if self.dark else "tf_logo_light.png"
        b64 = image_b64(logo_file)
        self.bg_els.append(f'<image x="{self.M}" y="24" width="248" height="32" href="{b64}"/>')

        # Vertical separator & programme title
        sep_col = PAL["TF_DARK_LINE"] if self.dark else "#C5D8CC"
        self.bg_els.append(f'<line x1="{self.M + 268}" y1="26" x2="{self.M + 268}" y2="54" stroke="{sep_col}" stroke-width="1.5"/>')

        tag_col = PAL["DARK_MUTED"] if self.dark else PAL["SLATE"]
        self.add_text(
            self.M + 284, 30, 600, 24,
            "TEFOS QUIZ  ·  NOTEBOOK QR PROGRAMME",
            size_px=13, weight="bold", color=tag_col, letter_spacing=1.6
        )

        # Right side: Section tag + Slide counter pill
        if self.count:
            pill_bg = PAL["TF_DARK_CARD"] if self.dark else PAL["TF_MINT"]
            pill_bdr = PAL["TF_DARK_LINE"] if self.dark else PAL["TF_MINT_BDR"]
            num_col = PAL["TF_LIME_BRT"] if self.dark else PAL["TF_EMERALD"]
            lbl_col = PAL["DARK_MUTED"] if self.dark else PAL["SLATE"]

            px_w = 290
            px_x = self.W - self.M - px_w
            self.rect(px_x, 21, px_w, 36, fill=pill_bg, stroke=pill_bdr, rx=18)
            self.add_text(
                px_x + 18, 29, 195, 20,
                section_label.upper(),
                size_px=11.5, weight="bold", color=lbl_col, letter_spacing=1.2
            )
            self.add_text(
                px_x + 210, 27, 66, 24,
                f"{self.slide_num:02d} / 22",
                size_px=13.5, weight="extrabold", color=num_col, align="right"
            )

        # Full-width Telecom Foundation Lime hairline rule (like Quetta before & after.pdf)
        line_col = "#1B5239" if self.dark else "#CBE3BA"
        self.bg_els.append(f'<line x1="{self.M}" y1="72" x2="{self.W - self.M}" y2="72" stroke="{line_col}" stroke-width="1.5"/>')

    def add_eyebrow(self, text_str, emoji="sparkles", x=None, y=92, color=None, bg_color=None):
        """Draws a rounded pill eyebrow with a Twemoji icon and uppercase tracking."""
        if x is None:
            x = self.M
        if color is None:
            color = PAL["TF_LIME_BRT"] if self.dark else PAL["TF_EMERALD"]
        if bg_color is None:
            bg_color = "rgba(124,192,64,0.14)" if self.dark else PAL["TF_LIME_PALE"]

        upper = text_str.upper()
        tw = measure_text_px(upper, "Poppins", "bold", 13.5, letter_spacing=2.0)
        pill_w = int(tw + (52 if emoji else 28))
        pill_h = 34
        self.bg_els.append(
            f'<rect x="{x}" y="{y}" width="{pill_w}" height="{pill_h}" rx="17" fill="{bg_color}"/>'
        )
        tx = x + 14
        if emoji:
            self.add_twemoji(emoji, x + 12, y + 6, 21)
            tx = x + 40
        self.add_text(
            tx, y + 8, pill_w - 20, 22,
            upper, size_px=13.5, weight="bold", color=color, letter_spacing=2.0
        )

    def add_title(self, title, highlight=None, sub=None, x=None, y=136, w=1728,
                  size_px=44, sub_size_px=19.5, highlight_color=None):
        """Adds the slide headline (with optional colored highlight phrase) and subtitle."""
        if x is None:
            x = self.M
        main_col = PAL["WHITE"] if self.dark else PAL["INK"]
        if highlight_color is None:
            highlight_color = PAL["TF_LIME_BRT"] if self.dark else PAL["TF_EMERALD"]
        sub_col = PAL["DARK_MUTED"] if self.dark else PAL["SLATE"]

        if highlight and highlight in title:
            parts = title.split(highlight, 1)
            runs = []
            if parts[0]:
                runs.append((parts[0], {"color": main_col, "weight": "extrabold", "size_px": size_px}))
            runs.append((highlight, {"color": highlight_color, "weight": "extrabold", "size_px": size_px}))
            if parts[1]:
                runs.append((parts[1], {"color": main_col, "weight": "extrabold", "size_px": size_px}))
        else:
            runs = [(title, {"color": main_col, "weight": "extrabold", "size_px": size_px})]

        used_h = self.add_rich_text(x, y, w, 110, [runs], line_height=1.12)
        if sub:
            sub_y = y + max(58, used_h + 6)
            self.add_text(x, sub_y, w, 60, sub, size_px=sub_size_px, weight="regular", color=sub_col, line_height=1.3)

    def add_footer(self, text_str, emoji="bulb"):
        y = self.H - 54
        line_col = "#1B5239" if self.dark else "#DCE6DF"
        self.bg_els.append(f'<line x1="{self.M}" y1="{y - 10}" x2="{self.W - self.M}" y2="{y - 10}" stroke="{line_col}" stroke-width="1"/>')
        tx = self.M
        if emoji:
            self.add_twemoji(emoji, self.M, y, 20)
            tx = self.M + 28
        col = PAL["DARK_MUTED"] if self.dark else PAL["SLATE"]
        self.add_text(tx, y + 1, self.W - 2 * self.M - 40, 26, text_str, size_px=14, weight="medium", color=col)

    def rect(self, x, y, w, h, fill="#FFFFFF", stroke=None, stroke_w=1.5, rx=16, shadow=False, opacity=1.0):
        flt = ' filter="url(#card_shadow)"' if shadow == "sm" else (' filter="url(#card_shadow_lg)"' if shadow else '')
        stk = f' stroke="{stroke}" stroke-width="{stroke_w}"' if stroke else ''
        op = f' opacity="{opacity}"' if opacity < 1.0 else ''
        self.bg_els.append(
            f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" '
            f'rx="{rx}" fill="{fill}"{stk}{flt}{op}/>'
        )

    def card(self, x, y, w, h, fill=None, stroke=None, rx=20, shadow="sm",
             top_stripe=None, left_stripe=None, stripe_w=7, bg_image=None, bg_opacity=1.0, overlay=0.42):
        if fill is None:
            fill = PAL["TF_DARK_CARD"] if self.dark else PAL["WHITE"]
        if stroke is None and not bg_image:
            stroke = PAL["TF_DARK_LINE"] if self.dark else PAL["BORDER"]

        self.rect(x, y, w, h, fill=fill, stroke=stroke, rx=rx, shadow=shadow)

        if bg_image:
            cid = self._next_id("cardimg")
            self.defs.append(f'<clipPath id="{cid}"><rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{rx}"/></clipPath>')
            b64 = image_b64(bg_image)
            self.bg_els.append(
                f'<image x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" '
                f'preserveAspectRatio="xMidYMid slice" clip-path="url(#{cid})" opacity="{bg_opacity}" href="{b64}"/>'
            )
            if overlay:
                self.bg_els.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{rx}" fill="#052316" opacity="{overlay}"/>')
            if stroke:
                self.rect(x, y, w, h, fill="none", stroke=stroke, rx=rx, shadow=False)

        if top_stripe:
            cid = self._next_id("topstripe")
            self.defs.append(f'<clipPath id="{cid}"><rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{rx}"/></clipPath>')
            self.bg_els.append(
                f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{stripe_w}" fill="{top_stripe}" clip-path="url(#{cid})"/>'
            )
        if left_stripe:
            cid = self._next_id("leftstripe")
            self.defs.append(f'<clipPath id="{cid}"><rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{rx}"/></clipPath>')
            self.bg_els.append(
                f'<rect x="{x:.1f}" y="{y:.1f}" width="{stripe_w}" height="{h:.1f}" fill="{left_stripe}" clip-path="url(#{cid})"/>'
            )

    def add_image(self, rel_path, x, y, w, h, rx=0, fit="contain", shadow=False):
        b64 = image_b64(rel_path)
        par = "xMidYMid meet" if fit == "contain" else "xMidYMid slice"
        flt = ' filter="url(#card_shadow_lg)"' if shadow else ''
        if rx > 0:
            cid = self._next_id("img")
            self.defs.append(f'<clipPath id="{cid}"><rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{rx}"/></clipPath>')
            self.bg_els.append(
                f'<g{flt}><image x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" '
                f'preserveAspectRatio="{par}" clip-path="url(#{cid})" href="{b64}"/></g>'
            )
        else:
            self.bg_els.append(
                f'<image x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" '
                f'preserveAspectRatio="{par}"{flt} href="{b64}"/>'
            )

    def add_twemoji(self, name, x, y, size=32):
        vb, inner = load_twemoji_svg(name)
        self.bg_els.append(
            f'<svg x="{x:.1f}" y="{y:.1f}" width="{size}" height="{size}" viewBox="{vb}">{inner}</svg>'
        )

    def add_lucide(self, name, x, y, size=28, color="#008450", stroke_width=2.2):
        vb, inner = load_lucide_svg(name, stroke_color=color, stroke_width=stroke_width)
        self.bg_els.append(
            f'<svg x="{x:.1f}" y="{y:.1f}" width="{size}" height="{size}" viewBox="{vb}">{inner}</svg>'
        )

    def icon_badge(self, x, y, size=52, lucide=None, emoji=None, bg="#EBF6EE",
                   border=None, icon_color="#008450", rx=14):
        """Draws a rounded icon container with either a Lucide vector icon or Twemoji."""
        self.rect(x, y, size, size, fill=bg, stroke=border, rx=rx)
        if emoji:
            esize = int(size * 0.56)
            off = (size - esize) / 2
            self.add_twemoji(emoji, x + off, y + off, esize)
        elif lucide:
            isize = int(size * 0.52)
            off = (size - isize) / 2
            self.add_lucide(lucide, x + off, y + off, isize, color=icon_color)

    def add_text(self, x, y, w, h, text_str, size_px=20, family="Poppins",
                 weight="regular", color=None, align="left", line_height=1.28,
                 letter_spacing=0.0, italic=False, v_align="top"):
        if color is None:
            color = PAL["WHITE"] if self.dark else PAL["INK"]
        paras = [
            [(line, {
                "size_px": size_px, "family": family, "weight": weight,
                "color": color, "letter_spacing": letter_spacing, "italic": italic
            })]
            for line in str(text_str).split("\n")
        ]
        return self.add_rich_text(
            x, y, w, h, paras, align=align, line_height=line_height,
            para_gap=int(size_px * 0.35), v_align=v_align
        )

    def add_rich_text(self, x, y, w, h, paragraphs, align="left", line_height=1.26,
                      para_gap=8, v_align="top"):
        """
        Wraps and renders multi-style paragraphs accurately using PIL font metrics,
        and records corresponding editable text for PowerPoint.
        Returns total rendered height in pixels.
        """
        default_col = PAL["WHITE"] if self.dark else PAL["INK"]
        wrapped_lines = []  # list of (line_runs, max_size_px, is_para_end)

        expanded = []
        for para in paragraphs:
            if isinstance(para, str):
                para = [(para, {})]
            cur = []
            for chunk_text, style in para:
                pieces = str(chunk_text).split("\n")
                for pi, piece in enumerate(pieces):
                    if pi > 0:
                        expanded.append(cur); cur = []
                    if piece:
                        cur.append((piece, style))
            expanded.append(cur)
        for para in expanded:
            # Tokenize runs into words preserving whitespace
            tokens = []
            for chunk_text, style in para:
                st = {
                    "size_px": style.get("size_px", 20),
                    "family": style.get("family", "Poppins"),
                    "weight": str(style.get("weight", "regular")).lower(),
                    "color": style.get("color", default_col),
                    "letter_spacing": style.get("letter_spacing", 0.0),
                    "italic": style.get("italic", False),
                }
                parts = re.split(r'(\s+)', str(chunk_text))
                for p in parts:
                    if p:
                        tokens.append((p, st))

            cur_line = []
            cur_w = 0.0
            for tok_str, st in tokens:
                if tok_str.isspace() and not cur_line:
                    continue
                tw = measure_text_px(tok_str, st["family"], st["weight"], st["size_px"], st["letter_spacing"])
                if cur_line and (cur_w + tw > w + 2.0) and not tok_str.isspace():
                    wrapped_lines.append((cur_line, max(r[1]["size_px"] for r in cur_line), False))
                    cur_line = [(tok_str.lstrip(), st)]
                    cur_w = measure_text_px(tok_str.lstrip(), st["family"], st["weight"], st["size_px"], st["letter_spacing"])
                else:
                    cur_line.append((tok_str, st))
                    cur_w += tw

            if cur_line:
                wrapped_lines.append((cur_line, max(r[1]["size_px"] for r in cur_line), True))

        # Compute total height
        total_h = 0.0
        for idx, (lruns, msize, is_end) in enumerate(wrapped_lines):
            total_h += msize * line_height
            if is_end and idx < len(wrapped_lines) - 1:
                total_h += para_gap

        curr_y = y
        if v_align == "middle" and h > total_h:
            curr_y = y + (h - total_h) / 2.0

        weight_num_map = {
            "regular": "400", "400": "400",
            "medium": "500", "500": "500",
            "semibold": "600", "600": "600",
            "bold": "700", "700": "700",
            "extrabold": "800", "800": "800",
        }

        for idx, (lruns, msize, is_end) in enumerate(wrapped_lines):
            # Strip trailing whitespace on line
            if lruns and lruns[-1][0].isspace():
                lruns = lruns[:-1]
            if not lruns:
                curr_y += msize * line_height
                continue

            baseline_y = curr_y + msize * 0.84
            anchor = "start"
            tx = x
            if align == "center":
                anchor = "middle"
                tx = x + w / 2.0
            elif align == "right":
                anchor = "end"
                tx = x + w

            tspans = []
            for r_txt, st in lruns:
                esc = html.escape(r_txt)
                fw = weight_num_map.get(st["weight"], "400")
                fs = ' font-style="italic"' if st.get("italic") else ''
                ls = f' letter-spacing="{st["letter_spacing"]}px"' if st.get("letter_spacing") else ''
                tspans.append(
                    f'<tspan font-family="{st["family"]}" font-weight="{fw}" '
                    f'font-size="{st["size_px"]:.1f}" fill="{st["color"]}"{fs}{ls}>{esc}</tspan>'
                )

            self.text_els.append(
                f'<text x="{tx:.1f}" y="{baseline_y:.1f}" text-anchor="{anchor}" xml:space="preserve">'
                + "".join(tspans) + '</text>'
            )

            curr_y += msize * line_height
            if is_end and idx < len(wrapped_lines) - 1:
                curr_y += para_gap

        # Record for PowerPoint editable layer
        self.pptx_texts.append({
            "x": x, "y": y, "w": w, "h": max(h, total_h + 8),
            "paragraphs": paragraphs, "align": align, "v_align": v_align,
            "default_col": default_col
        })
        return total_h

    def render_svg(self, include_text=True):
        body = "\n".join(self.bg_els)
        if include_text:
            body += "\n" + "\n".join(self.text_els)
        defs_str = "\n".join(self.defs)
        return (
            f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
            f'width="{self.W}" height="{self.H}" viewBox="0 0 {self.W} {self.H}">\n'
            f'<defs>{defs_str}</defs>\n'
            f'{body}\n</svg>'
        )

    def render_png_bytes(self, include_text=True):
        svg_str = self.render_svg(include_text=include_text)
        return resvg_py.svg_to_bytes(
            svg_string=svg_str,
            font_dirs=[FONT_DIR, DEJAVU_DIR]
        )


def export_deck(slides, pdf_paths, pptx_paths):
    """
    Exports the 22 SlideCanvas objects to:
      - High-resolution PNGs in tefos-quiz/slides/
      - High-resolution widescreen PDF in pdf_paths
      - Editable widescreen PPTX in pptx_paths
    """
    pdf_doc = pymupdf.open()
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for sc in slides:
        # 1. Full slide PNG (with embedded Poppins & Bebas Neue typography)
        full_png = sc.render_png_bytes(include_text=True)
        slide_png_path = os.path.join(SLIDES_OUT_DIR, f"slide_{sc.slide_num:02d}.png")
        with open(slide_png_path, "wb") as f:
            f.write(full_png)

        # Add page to PDF (960 x 540 pt = 13.333 x 7.5 in widescreen)
        page = pdf_doc.new_page(width=960, height=540)
        page.insert_image(pymupdf.Rect(0, 0, 960, 540), stream=full_png)

        # 2. PowerPoint slide: background art + cards + icons + logos (without text)
        #    topped with native editable PowerPoint text frames!
        bg_png = sc.render_png_bytes(include_text=False)
        bg_png_path = os.path.join(SLIDES_OUT_DIR, f"bg_{sc.slide_num:02d}.png")
        with open(bg_png_path, "wb") as f:
            f.write(bg_png)

        ps = prs.slides.add_slide(blank_layout)
        ps.shapes.add_picture(bg_png_path, Inches(0), Inches(0), Inches(13.333333), Inches(7.5))

        for item in sc.pptx_texts:
            ix = Inches(item["x"] / 144.0)
            iy = Inches(item["y"] / 144.0)
            iw = Inches(item["w"] / 144.0)
            ih = Inches(item["h"] / 144.0)
            tb = ps.shapes.add_textbox(ix, iy, iw, ih)
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = 0
            tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE if item["v_align"] == "middle" else MSO_ANCHOR.TOP

            align_map = {
                "left": PP_ALIGN.LEFT,
                "center": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT,
            }
            p_align = align_map.get(item["align"], PP_ALIGN.LEFT)

            for p_idx, para in enumerate(item["paragraphs"]):
                if isinstance(para, str):
                    para = [(para, {})]
                p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
                p.alignment = p_align
                for r_txt, st in para:
                    r = p.add_run()
                    r.text = str(r_txt)
                    f = r.font
                    f.name = st.get("family", "Poppins")
                    # 144 px = 1 inch = 72 pt -> 1 pt = 2.0 px
                    f.size = Pt(st.get("size_px", 20) / 2.0)
                    w_str = str(st.get("weight", "regular")).lower()
                    f.bold = w_str in ("semibold", "bold", "extrabold", "600", "700", "800")
                    f.italic = bool(st.get("italic", False))
                    col_hex = st.get("color", item["default_col"])
                    rgb = hex_to_rgb_tuple(col_hex)
                    f.color.rgb = RGBColor(*rgb)

    for p in pdf_paths:
        pdf_doc.save(p, deflate=True)
    pdf_doc.close()

    for p in pptx_paths:
        prs.save(p)

    # Clean up temporary bg_*.png files so slides/ only holds the 22 final slide PNGs
    for sc in slides:
        bgp = os.path.join(SLIDES_OUT_DIR, f"bg_{sc.slide_num:02d}.png")
        if os.path.exists(bgp):
            os.remove(bgp)
